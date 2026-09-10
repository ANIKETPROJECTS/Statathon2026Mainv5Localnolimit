from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "software_usage_steps.md")
OUTPUT = ROOT / (sys.argv[2] if len(sys.argv) > 2 else "AIRAVATA_DEA_Software_Usage_Steps.docx")


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_inline_runs(paragraph, text):
    token_pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    for token in token_pattern.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(token)


def add_body_paragraph(document, text, style=None):
    paragraph = document.add_paragraph(style=style)
    add_inline_runs(paragraph, text)
    return paragraph


def add_image(document, relative_path, alt_text):
    image_path = ROOT / relative_path
    if not image_path.exists():
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"[Screenshot unavailable: {relative_path}]")
        run.italic = True
        run.font.color.rgb = RGBColor(180, 60, 60)
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.45))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(alt_text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(90, 90, 90)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in [
        ("Title", 24, "1D4ED8"),
        ("Heading 1", 17, "1E3A8A"),
        ("Heading 2", 14, "1D4ED8"),
        ("Heading 3", 11, "334155"),
    ]:
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10)

    if "Screenshot Caption" not in [style.name for style in document.styles]:
        style = document.styles.add_style("Screenshot Caption", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Aptos"
        style.font.size = Pt(9)
        style.font.italic = True
        style.font.color.rgb = RGBColor(90, 90, 90)
        style.paragraph_format.space_after = Pt(8)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("AIRAVATA DEA — Software Usage Guide")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(110, 110, 110)


def create_document():
    document = Document()
    configure_document(document)
    document.core_properties.title = "AIRAVATA DEA — Software Usage Steps"
    document.core_properties.subject = "User guide for the AIRAVATA DEA application"
    document.core_properties.author = "AIRAVATA DEA"

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    first_heading = True
    index = 0

    while index < len(lines):
        line = lines[index].rstrip()

        if not line:
            index += 1
            continue

        if line == "---":
            index += 1
            continue

        linked_image_match = re.match(r"\[!\[([^\]]+)\]\(([^)]+)\)\]\([^)]+\)", line)
        plain_image_match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", line)
        image_match = linked_image_match or plain_image_match
        if image_match:
            image_path = image_match.group(2)
            alt_text = image_match.group(1)
            add_image(document, image_path, alt_text)
            index += 1
            continue

        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_runs(paragraph, line[2:])
            first_heading = False
            index += 1
            continue

        if line.startswith("## "):
            if not first_heading:
                document.add_page_break()
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline_runs(paragraph, line[3:])
            first_heading = False
            index += 1
            continue

        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline_runs(paragraph, line[4:])
            index += 1
            continue

        if line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.25)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline_runs(paragraph, line[2:])
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(60, 70, 90)
            index += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            text = re.sub(r"^\d+\.\s+", "", line)
            add_body_paragraph(document, text, style="List Number")
            index += 1
            continue

        if line.startswith("- "):
            text = line[2:]
            add_body_paragraph(document, text, style="List Bullet")
            index += 1
            continue

        if line.startswith("**Caption:**"):
            paragraph = document.add_paragraph(style="Screenshot Caption")
            add_inline_runs(paragraph, line)
            index += 1
            continue

        add_body_paragraph(document, line)
        index += 1

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    create_document()